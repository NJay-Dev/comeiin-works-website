"""Read the client workbook without modifying it. Rebuild catalogue and provenance."""

import json, pathlib, hashlib, copy, argparse
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument(
    "--workbook",
    type=pathlib.Path,
    default=ROOT / "source-files/Copy of Comeiin Works Product Review.xlsx",
)
SOURCE = parser.parse_args().workbook
wb = openpyxl.load_workbook(SOURCE, data_only=True)
sheet = wb["Products to review"]
headers = [c.value for c in sheet[5]]
references = {
    r[0]: {"name": r[1], "description": r[2], "page": r[3], "image": r[4]}
    for r in wb["Website reference"].iter_rows(min_row=6, max_row=39, values_only=True)
}
raw = []
products = []
exclusions = []
equipment = {1, 2, 7, 8, 9, 25}
liquid = {5, 19, 28}
accessories = {29, 30, 31, 32, 34}
generated = {"CW-029", "CW-030", "CW-031", "CW-034", "CW-023", "CW-001-SLIDES"}
for row in sheet.iter_rows(min_row=6, values_only=False):
    values = [c.value for c in row]
    if not values[0]:
        continue
    record = dict(zip(headers, values))
    record["sourceRow"] = row[0].row
    raw.append(record)
    ref = values[0]
    n = int(ref.split("-")[1])
    # Only actual client additions count as extra detail, not inherited reference links.
    has_details = any(values[i] for i in [5, 6, 8, 9, 10, 15, 16]) or ref == "CW-033"
    if values[1] == "Remove":
        exclusions.append(
            {
                "reference": ref,
                "reason": "Client action: Remove",
                "priority": values[17],
            }
        )
        continue
    if str(values[17]).lower() != "launch first" and not has_details:
        exclusions.append(
            {
                "reference": ref,
                "reason": "Not Launch first and no client-added technical details or image",
            }
        )
        continue
    cat = (
        "Equipment"
        if n in equipment
        else (
            "Liquid handling"
            if n in liquid
            else "Lab essentials" if n in accessories else "Consumables"
        )
    )
    p = {
        "id": ref,
        "reference": ref,
        "name": values[3],
        "sourceName": values[3],
        "category": cat,
        "description": str(values[7] or "").split("\nLaboratory page:")[0],
        "sourceListingType": values[2],
        "kind": "range",
        "specs": {},
        "variants": [],
        "pack": values[10] or "",
        "availability": values[13] or "On request",
        "priceBasis": values[12] or "Request quote",
        "sourceRow": row[0].row,
        "launchPriority": values[17] or "",
        "image": f"assets/products/{ref}-v1.jpg",
        "imageSource": references[ref]["image"],
        "imageOrigin": "Existing website mockup",
        "imageStatus": "Review required",
        "note": "",
    }
    if ref == "CW-001":
        p["category"] = "Equipment"
        p["pack"] = ""
        p["note"] = "Slide details from this row are assigned only to CW-001-SLIDES."
    elif ref == "CW-017":
        p.update(
            name="Universal specimen containers",
            kind="product",
            description="Sterile universal specimen containers for sample collection.",
            variants=["30 ml", "40 ml", "60 ml", "120 ml"],
            pack="500 pieces per box",
        )
        p["specs"] = {
            "Sterility": "Sterile",
            "Capacities": "30 ml, 40 ml, 60 ml and 120 ml",
            "Pack quantity": "500 pieces per box",
        }
    elif ref == "CW-023":
        p.update(
            name="Alcohol swabs",
            kind="product",
            description="Individually packed sterile antiseptic wipes.",
            pack="200 per box",
        )
        p["specs"] = {
            "Product type": "Alcohol swabs",
            "Sterility": "Sterile",
            "Pack quantity": "200 per box",
        }
    elif ref == "CW-033":
        p.update(
            name="Digital fridge / freezer thermometer",
            kind="product",
            category="Equipment",
            description="Digital thermometer for monitoring fridge and freezer temperatures.",
            pack="Each",
            image="assets/products/CW-033-client-v1.jpeg",
            imageSource=f"{SOURCE.name} — embedded image spanning P37:P39; visually identified as thermometer",
            imageOrigin="Client workbook image",
            imageStatus="Client supplied; final approval not recorded",
        )
        p["specs"] = {"Temperature range": "−50 °C to +70 °C", "Order unit": "Each"}
    if ref in generated:
        p.update(
            image=f"assets/products/{ref}-generated-v1.png",
            imageOrigin="AI-generated illustration",
            imageSource="",
            imageStatus="Generation pending; client review required",
        )
    products.append(p)
    if ref == "CW-001":
        slides = copy.deepcopy(p)
        slides.update(
            id="CW-001-SLIDES",
            name="Microscope slides",
            category="Consumables",
            description="Glass microscope slides supplied in a vacuum-sealed, moisture- and gas-impermeable liner.",
            kind="product",
            pack="50 pieces per box",
            specs={
                "Width": "26 mm",
                "Length": "76 mm",
                "Thickness": "1–1.2 mm",
                "Pack quantity": "50 pieces per box",
                "Packaging": "Vacuum-sealed, moisture- and gas-impermeable liner",
            },
            image="assets/products/CW-001-SLIDES-generated-v1.png",
            imageOrigin="AI-generated illustration",
            imageSource="",
            imageStatus="Generation pending; client review required",
            note="Separate slide listing derived from J6 and K6; not microscope specifications.",
        )
        products.append(slides)
for image in sheet._images:
    (ROOT / "dist/assets/products/CW-033-client-v1.jpeg").write_bytes(image._data())
(ROOT / "content/review-source.json").write_text(
    json.dumps(
        {
            "filename": SOURCE.name,
            "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            "sheet": sheet.title,
            "rows": raw,
            "exclusions": exclusions,
        },
        indent=2,
        ensure_ascii=False,
    )
)
(ROOT / "content/catalogue.json").write_text(
    json.dumps(products, indent=2, ensure_ascii=False)
)
(ROOT / "dist/catalogue.js").write_text(
    "window.COMEIIN_PRODUCTS = "
    + json.dumps(products, ensure_ascii=False, indent=2)
    + ";\n"
)
print(
    json.dumps(
        {
            "selectedSourceGroups": len(set(p["reference"] for p in products)),
            "catalogueEntries": len(products),
            "detailedProducts": sum(p["kind"] == "product" for p in products),
            "excluded": exclusions,
        },
        indent=2,
    )
)
